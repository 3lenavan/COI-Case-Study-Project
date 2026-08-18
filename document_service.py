from pathlib import Path
from docx import Document


# Folder where generated case study documents will be saved.
OUTPUT_FOLDER = Path("generated_case_studies")


# Save a generated case study draft as a Microsoft Word document.
def save_case_study_document(
    client_name: str,
    case_study_draft: str,
) -> str:

    # Create the output folder if it does not already exist.
    OUTPUT_FOLDER.mkdir(exist_ok=True)

    # Make the client name safer to use as a file name.
    safe_client_name = client_name.replace(" ", "_")

    # Create the Word document file path.
    file_path = OUTPUT_FOLDER / f"{safe_client_name}_case_study.docx"

    # Create a new Microsoft Word document.
    document = Document()

    # Add a title to the document.
    document.add_heading(
        f"{client_name} Case Study",
        level=0,
    )

    # Add the AI-generated case study draft.
    document.add_paragraph(case_study_draft)

    # Save the Word document.
    document.save(file_path)

    # Return the location of the saved document.
    return str(file_path)


# Test Word document creation when this file is run directly.
if __name__ == "__main__":
    test_draft = """
Client Overview

Patriot Family Insurance is the client for the Nashua Office Renovation project.

Project Challenge

The client needs new furniture and workspace solutions.

COI Solution

COI proposed furniture and workspace solutions.

Project Results

Project results are not yet available.
"""

    file_path = save_case_study_document(
        client_name="Patriot Family Insurance",
        case_study_draft=test_draft,
    )

    print("\n========== WORD DOCUMENT TEST ==========")
    print(f"Case study saved to: {file_path}")